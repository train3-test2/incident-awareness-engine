"""Windows-only native Word verification. Run with an ephemeral pywin32 environment.

Opens in Word print layout, captures native page metafiles, edits a separate copy,
and reopens that copy. Does not generate, print, or convert a PDF.
"""

import argparse
import hashlib
import json
from pathlib import Path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-file", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--reject-isolated-response", action="store_true")
    parser.add_argument("--inspect-paragraphs", action="store_true")
    args = parser.parse_args()
    import pythoncom
    import win32com.client

    source = args.input_file.resolve(strict=True)
    from docx import Document
    from docx.oxml.ns import qn

    package = Document(source)
    response_length = 0
    for row in package.tables[1]._tbl.findall(qn("w:tr")):
        if "기업 조치사항" in "".join(t.text or "" for t in row.iter(qn("w:t"))):
            response = row.findall(qn("w:tc"))[-1].find(qn("w:p"))
            response_length = len("".join(t.text or "" for t in response.iter(qn("w:t"))))
    output = args.output_dir.resolve()
    if not output.is_relative_to(Path(__file__).resolve().parents[2] / "build"):
        parser.error("Verification outputs must be inside the repository build directory")
    output.mkdir(parents=True, exist_ok=True)
    edited = output / "word-edit-check.docx"
    if source == edited:
        parser.error("Input must be different from output copies")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    pythoncom.CoInitialize()
    application = document = None
    owned = False
    pid = None

    def word_ids():
        return {
            int(p.ProcessId)
            for p in win32com.client.GetObject("winmgmts:").ExecQuery(
                "SELECT ProcessId FROM Win32_Process WHERE Name='WINWORD.EXE'"
            )
        }

    def progress(stage):
        (output / "word-progress.json").write_text(
            json.dumps({"stage": stage, "owned_word_pid": pid}), encoding="utf8"
        )

    try:
        existing = word_ids()
        application = win32com.client.DispatchEx("Word.Application")
        _ = application.Documents.Count
        created = word_ids() - existing
        if len(created) != 1:
            raise RuntimeError(
                "Cannot establish unique ownership of Word; existing instances preserved"
            )
        pid = created.pop()
        owned = True
        application.Visible = False
        application.DisplayAlerts = 0
        application.AutomationSecurity = 3
        progress("open-readonly")
        document = application.Documents.Open(str(source), False, True, False)
        document.ActiveWindow.View.Type = 3  # wdPrintView; real Word pagination, no PDF.
        document.Repaginate()
        pages = document.ComputeStatistics(2)
        tables = document.Tables.Count
        (output / "word-statistics.json").write_text(
            json.dumps({"pages": pages, "tables": tables}), encoding="utf8"
        )
        progress("edit-copy")
        input_cell = document.Content.Duplicate
        if not input_cell.Find.Execute(FindText="- 기업명 :"):
            raise RuntimeError("Company input cell not found for native editing")
        document.Range(input_cell.End, input_cell.End).InsertBefore(" [WORD EDIT CHECK ONLY]")
        document.SaveAs2(str(edited), 16)
        document.Close(0)
        document = application.Documents.Open(str(edited), False, True, False)
        if "[WORD EDIT CHECK ONLY]" not in document.Content.Text:
            raise RuntimeError("Native Word edit was not persisted")
        document.Close(0)
        document = application.Documents.Open(str(source), False, True, False)
        document.ActiveWindow.View.Type = 3
        document.Repaginate()
        progress("capture-print-layout")
        from PIL import Image

        layout = []
        if args.inspect_paragraphs:
            for index, paragraph in enumerate(document.Tables(2).Range.Paragraphs):
                span = paragraph.Range
                layout.append(
                    {
                        "index": index,
                        "characters": len(span.Text),
                        "page": span.Information(3),
                        "font_size": span.Font.Size,
                        "lines": span.ComputeStatistics(1),
                        "line_spacing": paragraph.Format.LineSpacing,
                        "space_before": paragraph.Format.SpaceBefore,
                        "space_after": paragraph.Format.SpaceAfter,
                    }
                )
            (output / "paragraph-layout.json").write_text(
                json.dumps(layout, indent=2), encoding="utf8"
            )
        document.Repaginate()
        pages = document.ComputeStatistics(2)
        (output / "word-statistics.json").write_text(
            json.dumps({"pages": pages, "tables": tables}), encoding="utf8"
        )
        clean = lambda value: "".join(c for c in value if ord(c) >= 32 and not c.isspace())
        native_body = clean(document.Content.Text)
        for paragraph in package._element.body.iter(qn("w:p")):
            expected = clean("".join(t.text or "" for t in paragraph.iter(qn("w:t"))))
            if expected and expected not in native_body:
                raise RuntimeError("Word omitted source document text")
        footer_expected = "".join(p.text for p in package.sections[0].footer.paragraphs)
        if footer_expected and clean(footer_expected) not in clean(
            document.Sections(1).Footers(1).Range.Text
        ):
            raise RuntimeError("Word footer differs from the saved source document")
        response_label = document.Content.Duplicate
        response_page = pages
        if response_label.Find.Execute(FindText="기업 조치사항"):
            response_page = document.Range(response_label.Start, response_label.Start).Information(
                3
            )
        page_results = []
        for number in range(1, pages + 1):
            start = document.GoTo(1, 1, number).Start
            end = document.GoTo(1, 1, number + 1).Start if number < pages else document.Content.End
            content = document.Range(start, end).Text
            metafile = output / f"page-{number}.emf"
            bits = document.ActiveWindow.Panes(1).Pages(number).EnhMetaFileBits
            metafile.write_bytes(bytes(bits))
            with Image.open(metafile) as picture:
                picture.load()
                picture.save(output / f"page-{number}.png")
            compact = clean(content)
            isolated = (
                response_length < 250
                and number >= response_page
                and len(compact) < 350
                and "피해IP" not in compact
                and "피해도메인" not in compact
            )
            page_results.append(
                {
                    "page": number,
                    "body_characters": len(compact),
                    "isolated_short_response": isolated,
                    "image": str(output / f"page-{number}.png"),
                }
            )
        if (
            args.reject_isolated_response
            and pages > 2
            and page_results[-1]["isolated_short_response"]
        ):
            raise RuntimeError("Last print-layout page contains an isolated short response row")
        if before != hashlib.sha256(source.read_bytes()).hexdigest():
            raise RuntimeError("Original generated file changed")
        result = {
            "pages": pages,
            "tables": tables,
            "word_version": application.Version,
            "source_sha256": before,
            "original_unchanged": True,
            "opened_read_only": True,
            "edited_copy_saved_and_reopened": True,
            "edited_input_cell": True,
            "print_layout": True,
            "all_source_body_paragraphs_present": True,
            "native_footer_matches_source": bool(footer_expected),
            "page_layouts": page_results,
            "last_page_isolated_short_response": page_results[-1]["isolated_short_response"],
            "response_starts_on_page": response_page,
            "pdf_generated": False,
            "edited_copy": str(edited),
            "owned_word_process": pid,
        }
        (output / "word-validation.json").write_text(json.dumps(result, indent=2), encoding="utf8")
        progress("complete")
        print(json.dumps(result))
    finally:
        if document is not None:
            document.Close(0)
        if application is not None and owned:
            application.Quit(0)
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    main()
